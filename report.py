import json
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

# 预埋缺陷"答案清单"：用来计算缺陷发现率
KNOWN_BUGS = {
    "BUG1 密码长度16被误拒": lambda r: r["path"] == "/login"
        and isinstance(r["body"], dict)
        and isinstance(r["body"].get("password"), str)
        and len(r["body"]["password"]) == 16,
    "BUG2 user_id=0未拦截": lambda r: r["path"] == "/users/0",
    "BUG3 下单数量10被误拒": lambda r: r["path"] == "/orders"
        and isinstance(r["body"], dict)
        and r["body"].get("quantity") == 10,
    "BUG4 size=0未拦截": lambda r: r["path"].startswith("/products")
        and "size=0" in r["path"],
}


def generate_report(results, path="测试报告.xlsx", show_bugs=True):
    """根据执行结果生成 Excel 报告（汇总 + 用例明细）。"""
    total = len(results)
    passed = sum(1 for r in results if r["result"] == "通过")
    failed_cases = [r for r in results if r["result"] != "通过"]
    executable = sum(1 for r in results if r["result"] != "错误")
    defects = [r for r in failed_cases if r.get("category") == "接口缺陷"]
    unreasonable = [r for r in failed_cases
                    if r.get("category") == "用例预期不合理"]

    found_bugs = {
        name: any(r["result"] == "失败" and check(r) for r in results)
        for name, check in KNOWN_BUGS.items()
    }
    found_count = sum(found_bugs.values())

    wb = Workbook()

    # 汇总表
    ws = wb.active
    ws.title = "汇总"
    rows = [
        ("用例总数", total),
        ("通过", passed),
        ("未通过", len(failed_cases)),
        ("用例可执行率", f"{executable / total:.0%}"),
        ("接口缺陷（AI归因）", len(defects)),
        ("用例预期不合理（AI归因）", len(unreasonable)),
        ("用例预期不合理占比", f"{len(unreasonable) / total:.0%}"),
        ("预埋缺陷总数", len(KNOWN_BUGS)),
        ("预埋缺陷发现数", found_count),
        ("预埋缺陷发现率", f"{found_count / len(KNOWN_BUGS):.0%}"),
    ]
    if not show_bugs:
        rows = rows[:7]  # 自定义要求时，不显示预埋缺陷相关统计
    for row in rows:
        ws.append(row)
    if show_bugs:
        ws.append(())
        ws.append(("预埋缺陷明细", "是否发现"))
        for name, ok in found_bugs.items():
            ws.append((name, "已发现" if ok else "未发现"))
    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 14

    # 用例明细表
    ws2 = wb.create_sheet("用例明细")
    headers = ["用例编号", "用例名称", "请求方法", "路径", "请求体",
               "预期状态码", "实际状态码", "结果", "归因类别", "原因", "建议"]
    ws2.append(headers)
    for cell in ws2[1]:
        cell.font = Font(bold=True)

    red = PatternFill("solid", fgColor="FFC7CE")
    green = PatternFill("solid", fgColor="C6EFCE")
    for r in results:
        ws2.append([
            r["id"], r["name"], r["method"], r["path"],
            json.dumps(r["body"], ensure_ascii=False) if r["body"] else "",
            r["expected_status"], r["actual_status"], r["result"],
            r.get("category", ""), r.get("reason", ""), r.get("suggestion", ""),
        ])
        ws2.cell(row=ws2.max_row, column=8).fill = (
            green if r["result"] == "通过" else red)

    for col, width in zip("ABCDEFGHIJK",
                          [10, 28, 10, 28, 40, 12, 12, 8, 16, 50, 40]):
        ws2.column_dimensions[col].width = width

    wb.save(path)
    if show_bugs:
        print(f"已生成 {path}，预埋缺陷发现率：{found_count}/{len(KNOWN_BUGS)}")
    else:
        print(f"已生成 {path}")